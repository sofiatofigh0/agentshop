"""The three renders of a resume, and what a parser makes of each.

The upload copies — single-column PDF and .docx — are tested for the things a
screening system trips on: reading order, text-only bullets, contact details in
the body, no tables. The two-column designed copy is tested for the opposite:
that it really does scramble when read as one stream, which is the reason it
is not the upload copy. If that test ever starts failing, the layout changed
and the advice in the UI should be revisited.
"""

import os
import re
import tempfile
import unittest

import ats
import documents

RESUME = """# Jane Example
**Senior Product Manager**
New York, NY · jane@example.com · 555-0100 · [portfolio](https://jane.example.com/)

## Summary
Senior Product Manager with four years shipping LLM products in regulated finance.

## Skills
- LLM evaluation
- SQL and analytics
- API integrations

## Work Experience
### AI Product Manager, Senior Associate — JPMorgan Chase | March 2026 - Present
- **Advisor summarization** — Shipped an LLM pipeline that cut meeting prep by ~40%.
- **Eval framework** — Built the rubric and reporting behind every model release.

### Product Manager — Axial | October 2022 - November 2025
- **Partner APIs** — Integrated two supply partners, expanding deal inventory.

## Education
**Full-Stack Coding Bootcamp**
Flatiron School · 2020

**BA, Economics**
Columbia University · September 2015 - May 2019
"""


class UploadPdf(unittest.TestCase):
    def setUp(self):
        self.html = documents.render_resume_ats_html(RESUME)
        self.css = documents.ATS_CSS(9.6, "0.5in")

    def test_no_layout_a_parser_can_misread(self):
        self.assertNotIn("<table", self.html)
        for rule in ("display: table", "display:table", "float", "column", "grid", "flex"):
            self.assertNotIn(rule, self.css, rule)

    def test_sections_come_out_in_the_order_they_were_written(self):
        headings = re.findall(r"<h2>(.*?)</h2>", self.html)
        self.assertEqual(headings, ["Summary", "Skills", "Work Experience", "Education"])

    def test_bullets_are_text_not_list_markers(self):
        """A drawn marker extracts as a stray character somewhere else on the page."""
        self.assertIn("list-style: none", self.css)
        self.assertIn("<li>• <strong>Advisor summarization</strong>", self.html)

    def test_skills_are_one_comma_separated_line(self):
        self.assertIn("<p>LLM evaluation, SQL and analytics, API integrations</p>", self.html)

    def test_a_link_shows_its_address(self):
        """A parser stores the text; a label hides the URL from it."""
        self.assertIn(">jane.example.com</a>", self.html)
        self.assertNotIn(">portfolio</a>", self.html)

    def test_role_line_keeps_its_dates_on_the_same_line(self):
        self.assertIn("<strong>Product Manager — Axial</strong>"
                      '<span class="jobdate"> | October 2022 - November 2025</span>', self.html)


class ReadingOrder(unittest.TestCase):
    """The claim that two columns scramble, checked on this project's own layouts."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.results = {}
        for style in ("resume", "resume_designed"):
            path = os.path.join(cls.tmp.name, f"{style}.pdf")
            pages, _ = documents.fit_pdf(RESUME, path, style)
            cls.results[style] = (pages, ats.reading_order(path, RESUME), path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_the_upload_copy_reads_as_one_stream(self):
        pages, order, _ = self.results["resume"]
        self.assertEqual(pages, 1)
        self.assertTrue(order["ok"], order)
        self.assertEqual(order["fused_rows"], 0)

    def test_the_designed_copy_scrambles_which_is_why_it_is_not_uploaded(self):
        pages, order, _ = self.results["resume_designed"]
        self.assertEqual(pages, 1)
        self.assertFalse(order["in_order"])      # the sidebar's Education jumps ahead
        self.assertGreater(order["fused_rows"], 0)  # rows fuse the two columns

    def test_every_term_survives_extraction_from_the_upload_copy(self):
        _, _, path = self.results["resume"]
        terms = [{"term": t, "aliases": []} for t in
                 ("advisor summarization", "jpmorgan chase", "flatiron school",
                  "march 2026", "september 2015")]
        self.assertTrue(ats.pdf_text_check(path, terms)["ok"])


class WordLibrary(unittest.TestCase):
    def test_availability_is_detected(self):
        import sys
        from unittest import mock
        self.assertTrue(documents.docx_available())
        with mock.patch.dict(sys.modules, {"docx": None}):
            self.assertFalse(documents.docx_available())

    def test_the_library_is_a_declared_requirement(self):
        here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        with open(os.path.join(here, "requirements.txt")) as handle:
            self.assertIn("python-docx", handle.read().split())


class WordCopy(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from docx import Document
        cls.tmp = tempfile.TemporaryDirectory()
        path = os.path.join(cls.tmp.name, "resume.docx")
        documents.write_resume_docx(RESUME, path, pt=9.6)
        cls.doc = Document(path)
        cls.paragraphs = [(p.style.name, p.text, [r.text for r in p.runs if r.bold])
                          for p in cls.doc.paragraphs]

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_nothing_lives_in_the_header_or_footer(self):
        """Some parsers skip them, and would lose whoever the resume belongs to."""
        section = self.doc.sections[0]
        self.assertEqual("".join(p.text for p in section.header.paragraphs), "")
        self.assertEqual("".join(p.text for p in section.footer.paragraphs), "")

    def test_no_tables_or_shapes(self):
        self.assertEqual(len(self.doc.tables), 0)
        self.assertEqual(len(self.doc.inline_shapes), 0)

    def test_name_title_and_contact_open_the_body(self):
        texts = [t for _, t, _ in self.paragraphs[:3]]
        self.assertEqual(texts[0], "Jane Example")
        self.assertEqual(texts[1], "Senior Product Manager")
        self.assertIn("jane@example.com", texts[2])
        self.assertIn("jane.example.com", texts[2])   # the address, not "portfolio"

    def test_sections_are_real_headings_in_order(self):
        headings = [t for style, t, _ in self.paragraphs if style == "Heading 1"]
        self.assertEqual(headings, ["Summary", "Skills", "Work Experience", "Education"])

    def test_role_line_is_one_line_with_a_tab_before_the_dates(self):
        roles = [t for _, t, _ in self.paragraphs if "\t" in t]
        self.assertEqual(roles[0],
                         "AI Product Manager, Senior Associate — JPMorgan Chase\tMarch 2026 - Present")

    def test_bullets_use_the_list_style_and_keep_the_bold_label(self):
        bullets = [(t, bold) for style, t, bold in self.paragraphs if style == "List Bullet"]
        self.assertEqual(len(bullets), 3)
        self.assertEqual(bullets[0][1], ["Advisor summarization"])

    def test_letter_size_and_body_font(self):
        section = self.doc.sections[0]
        self.assertAlmostEqual(section.page_width.inches, 8.5)
        self.assertAlmostEqual(section.page_height.inches, 11)
        self.assertEqual(self.doc.styles["Normal"].font.name, "Arial")


if __name__ == "__main__":
    unittest.main()
