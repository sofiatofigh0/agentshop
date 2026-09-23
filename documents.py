"""
PDF rendering.

The model writes markdown; this turns it into the finished document. Keeping
those separate matters — the writing prompts stay about content, and how a
resume looks is a styling decision made once here rather than negotiated with a
model on every run.

The resume is rendered three ways from the same markdown:

    ATS_CSS       tailored_resume.pdf — one column, no tables, bullets as text.
                  The copy to upload. A screening system reads a page as one
                  stream, and this layout is that stream.
    write_resume_docx
                  tailored_resume.docx — the same single column as a Word
                  document. The format every applicant tracking system parses
                  without edge cases; upload this unless a form asks for PDF.
    DOCUMENT_CSS  resume_designed.pdf — the two-column design. For people:
                  a referral, an email to a hiring manager, a printout. Not for
                  upload, because a parser reads two columns as one and fuses
                  them (tests/test_documents.py shows it on this very layout).

Plus LETTER_CSS for the cover letter, and REPORT_CSS for the internal working
documents — evidence map, factuality review, strategy, ATS report — which are
denser and allowed wide tables.
"""

import re

import markdown
from pypdf import PdfReader
from weasyprint import CSS, HTML

# --------------------------------------------------------------------------
# Shared page setup
# --------------------------------------------------------------------------

BASE = """
@page { size: Letter; margin: 0.5in 0.55in; }
* { box-sizing: border-box; }
body { margin: 0; color: #16181d; }
a { color: inherit; text-decoration: none; }
strong { font-weight: 600; }
hr { display: none; }
"""

# --------------------------------------------------------------------------
# Resume and cover letter
# --------------------------------------------------------------------------

# --------------------------------------------------------------------------
# Resume — two-column template
#
# Geometry and colour sampled from the supplied template PDF: Letter, 0.6in
# margins, header rule at 1.38in, sidebar/main divider at 2.77in (a 30/70
# split), accent #ae1800, text #201e1d, secondary #696767, rules #d8d7d7 on a
# #f3f2f2 ground.
#
# Everything is in em so the whole sheet still scales from one number, which is
# what lets fit_pdf() land it on one page without cutting content.
# --------------------------------------------------------------------------

ACCENT, INK, MUTED, RULE, GROUND = "#ae1800", "#201e1d", "#696767", "#d8d7d7", "#f3f2f2"

# Which sections live in the narrow left column. Everything else goes right.
SIDEBAR_SECTIONS = {"profile", "summary", "skills", "education", "recognition",
                    "languages", "certifications", "awards"}


def DOCUMENT_CSS(pt, margin):
    """Resume styling at a given body size."""
    return f"""
@page {{ size: Letter; margin: {margin}; background: {GROUND}; }}
* {{ box-sizing: border-box; }}
body {{ margin: 0; background: {GROUND}; color: {INK};
       font-family: "Helvetica Neue", Helvetica, Arial, sans-serif;
       font-size: {pt}pt; line-height: 1.4; }}
a {{ color: {ACCENT}; text-decoration: none; }}
strong {{ font-weight: 700; }}

/* ---- header ---- */
.head {{ display: table; width: 100%; }}
.head .who, .head .contact {{ display: table-cell; vertical-align: bottom; }}
.head .who {{ width: 48%; }}
.head .contact {{ text-align: right; font-size: 0.84em; color: {MUTED}; line-height: 1.55; }}
.name {{ font-size: 2.5em; font-weight: 700; letter-spacing: -0.015em;
        margin: 0; color: {INK}; white-space: nowrap; }}
.title {{ font-size: 1.02em; font-weight: 600; color: {ACCENT}; margin: 0.15em 0 0; }}
.rule {{ border-bottom: 1.1pt solid {INK}; margin: 0.75em 0 1.1em; }}

/* ---- two columns ---- */
.body {{ display: table; width: 100%; }}
.side, .main {{ display: table-cell; vertical-align: top; }}
.side {{ width: 30%; padding-right: 0.9em; }}
.main {{ border-left: 0.7pt solid {RULE}; padding-left: 1.15em; }}

/* ---- section headings ---- */
h2 {{ font-size: 0.79em; font-weight: 700; letter-spacing: 0.14em;
     text-transform: uppercase; color: {ACCENT};
     margin: 0 0 0.55em; page-break-after: avoid; }}
.side section, .main section {{ margin-bottom: 1.35em; }}
.side section:last-child, .main section:last-child {{ margin-bottom: 0; }}

/* ---- job entries ---- */
.job {{ margin-bottom: 0.95em; page-break-inside: avoid; }}
.job + .job {{ border-top: 0.7pt solid {RULE}; padding-top: 0.85em; }}
.jobhead {{ display: table; width: 100%; margin-bottom: 0.4em; }}
.jobtitle, .jobdate {{ display: table-cell; vertical-align: baseline; }}
.jobtitle {{ font-size: 1.06em; font-weight: 700; color: {INK}; line-height: 1.3; }}
.jobdate {{ text-align: right; white-space: nowrap; color: {MUTED};
           font-size: 0.88em; padding-left: 0.6em; }}
.jobnote {{ color: {MUTED}; font-size: 0.88em; margin: -0.2em 0 0.35em; }}

/* ---- text ---- */
p {{ margin: 0 0 0.5em; }}
p:last-child {{ margin-bottom: 0; }}
ul {{ margin: 0; padding-left: 1.05em; list-style-type: disc; }}
li {{ margin-bottom: 0.34em; padding-left: 0.15em; }}
li:last-child {{ margin-bottom: 0; }}
li::marker {{ color: {INK}; }}
.side li {{ margin-bottom: 0.28em; }}
.side p {{ color: {INK}; }}
.muted {{ color: {MUTED}; }}
"""


def ATS_CSS(pt, margin):
    """The upload copy: one column, in document order, nothing a parser can
    misread.

    Every rule here is about what survives text extraction. No tables, no
    floats and no columns, so the reading order is the visual order. Bullets
    are text rather than list markers, because a marker is drawn separately
    from its line and extracts as a stray character somewhere else on the
    page. Headings keep the case they were written in. Contact details sit in
    the body — nothing uses the page margins, which some parsers skip.
    """
    return f"""
@page {{ size: Letter; margin: {margin}; }}
* {{ box-sizing: border-box; }}
body {{ margin: 0; color: #111111;
       font-family: Arial, "Helvetica Neue", Helvetica, sans-serif;
       font-size: {pt}pt; line-height: 1.38; }}
a {{ color: inherit; text-decoration: none; }}
strong {{ font-weight: 700; }}

.name {{ font-size: 2.1em; font-weight: 700; margin: 0; }}
.title {{ font-size: 1.12em; font-weight: 700; margin: 0.1em 0 0.15em; }}
.contact {{ margin: 0 0 0.3em; }}

h2 {{ font-size: 1.05em; font-weight: 700; margin: 0.95em 0 0.4em;
     padding-bottom: 0.15em; border-bottom: 0.8pt solid #111111;
     page-break-after: avoid; }}

.job {{ margin-bottom: 0.7em; page-break-inside: avoid; }}
.jobhead {{ margin: 0 0 0.2em; }}
.jobdate {{ font-weight: 400; }}
.jobnote {{ font-style: italic; margin: 0 0 0.2em; }}

p {{ margin: 0 0 0.35em; }}
ul {{ list-style: none; margin: 0; padding: 0; }}
li {{ margin: 0 0 0.22em; padding-left: 1em; text-indent: -0.75em; }}
"""


def LETTER_CSS(pt, margin):
    """Cover letter styling at a given body size."""
    return BASE + f"""
@page {{ size: Letter; margin: {margin}; }}
body {{ font-family: Georgia, "Times New Roman", serif;
       font-size: {pt}pt; line-height: 1.58; }}
h1 {{ font-family: "Helvetica Neue", Helvetica, Arial, sans-serif;
     font-size: 1.6em; font-weight: 600; letter-spacing: 0.01em; margin: 0 0 0.25em; }}
h1 + p {{ font-family: "Helvetica Neue", Helvetica, Arial, sans-serif;
         font-size: 0.81em; color: #55595f; margin: 0 0 1.7em;
         padding-bottom: 0.85em; border-bottom: 0.6pt solid #cfd3d8; }}
p {{ margin: 0 0 0.95em; text-align: left; }}
p:last-child {{ margin-top: 0.35em; }}
"""


REPORT_CSS = BASE + """
@page { size: Letter landscape; margin: 0.45in; }
body { font-family: "Helvetica Neue", Helvetica, Arial, sans-serif;
       font-size: 8.4pt; line-height: 1.42; }
h1 { font-size: 15pt; font-weight: 600; margin: 0 0 8pt; }
h2 { font-size: 9.5pt; font-weight: 700; letter-spacing: 0.09em;
     text-transform: uppercase; color: #6a6f76; margin: 12pt 0 5pt;
     padding-bottom: 3pt; border-bottom: 0.6pt solid #dfe3e7; page-break-after: avoid; }
h3 { font-size: 9.4pt; font-weight: 600; margin: 8pt 0 3pt; }
p { margin: 0 0 5pt; }
ul { margin: 3pt 0 6pt; padding-left: 13pt; }
li { margin-bottom: 2.6pt; }
table { border-collapse: collapse; width: 100%; margin: 5pt 0 9pt;
        font-size: 7.5pt; page-break-inside: auto; }
th { background: #f2f4f6; text-align: left; font-size: 6.8pt; font-weight: 700;
     letter-spacing: 0.07em; text-transform: uppercase; color: #55595f;
     padding: 4pt 5pt; border: 0.5pt solid #d8dce1; }
td { padding: 4pt 5pt; border: 0.5pt solid #e3e7eb; vertical-align: top; }
tr { page-break-inside: avoid; }
code { font-family: "SF Mono", Menlo, Consolas, monospace; font-size: 7.4pt; }
"""

# Density ladders, loosest first. fit_pdf walks down until the document lands
# on one page. It never removes content — it only tightens the setting.
_RESUME_LADDER = [(9.6, "0.55in 0.6in"), (9.3, "0.5in 0.55in"), (9.0, "0.5in 0.5in"),
                  (8.7, "0.45in 0.5in"), (8.4, "0.42in 0.45in"), (8.1, "0.4in 0.45in"),
                  (7.8, "0.38in 0.42in"), (7.5, "0.35in 0.4in")]

LADDERS = {
    "resume": _RESUME_LADDER,
    "resume_designed": _RESUME_LADDER,
    "letter": [(10.8, "0.9in 0.95in"), (10.5, "0.85in 0.9in"), (10.2, "0.8in 0.85in"),
               (9.9, "0.75in 0.8in"), (9.6, "0.7in 0.75in"), (9.3, "0.65in 0.7in"),
               (9.0, "0.6in 0.7in")],
}
BUILDERS = {"resume": lambda pt, margin: ATS_CSS(pt, margin),
            "resume_designed": DOCUMENT_CSS, "letter": LETTER_CSS}


# --------------------------------------------------------------------------
# Resume markdown -> two-column HTML
#
# The model writes a flat document; the template is two columns. Rather than
# ask the model to emit HTML or guess at layout, it writes a known set of
# sections and this router decides which column each one belongs in. Layout
# stays a decision made once, in code.
# --------------------------------------------------------------------------

def _visible_url(href: str) -> str:
    """A URL as a reader would type it: no scheme, no trailing slash."""
    return re.sub(r"^https?://(www\.)?", "", href).rstrip("/")


def _inline(text: str, show_urls: bool = False) -> str:
    """Escape, then honour **bold** and [label](href).

    With `show_urls`, a link's visible text is its address rather than its
    label. A parser reads the text, not the link behind it, so a portfolio
    link labelled "portfolio" gives it nothing to store.
    """
    out = (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
    if show_urls:
        out = re.sub(r"\[([^\]]+)\]\(([^)]+)\)",
                     lambda m: f'<a href="{m.group(2)}">{_visible_url(m.group(2))}</a>', out)
    else:
        out = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', out)
    return re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", out)


def _parse_resume(md: str) -> dict:
    """Pull the resume apart into a header and a list of sections."""
    doc = {"name": "", "title": "", "contact": "", "sections": []}
    section = None
    job = None

    for raw in md.splitlines():
        line = raw.strip()
        if not line:
            continue

        if line.startswith("# "):
            doc["name"] = line[2:].strip()
            continue

        if line.startswith("## "):
            section = {"heading": line[3:].strip(), "blocks": []}
            doc["sections"].append(section)
            job = None
            continue

        # Before the first section: the title line, then the contact line.
        if section is None:
            stripped = line.strip("*")
            if line.startswith("**") and not doc["title"]:
                doc["title"] = stripped
            elif not doc["contact"]:
                doc["contact"] = line
            continue

        if line.startswith("### "):
            head = line[4:].strip()
            title, _, dates = head.rpartition("|")
            job = {"kind": "job", "title": (title or head).strip(),
                   "dates": dates.strip() if title else "", "note": "", "items": []}
            section["blocks"].append(job)
            continue

        if line.startswith(("- ", "* ")):
            item = line[2:].strip()
            if job is not None:
                job["items"].append(item)
            else:
                if not section["blocks"] or section["blocks"][-1]["kind"] != "list":
                    section["blocks"].append({"kind": "list", "items": []})
                section["blocks"][-1]["items"].append(item)
            continue

        # Plain prose. Directly under a job head it is that role's one-liner.
        if job is not None and not job["items"] and not job["note"]:
            job["note"] = line
        else:
            section["blocks"].append({"kind": "para", "text": line})

    return doc


def _render_blocks(blocks: list) -> str:
    html = []
    for block in blocks:
        if block["kind"] == "job":
            head = (f'<div class="jobhead"><span class="jobtitle">{_inline(block["title"])}</span>'
                    f'<span class="jobdate">{_inline(block["dates"])}</span></div>')
            note = f'<p class="jobnote">{_inline(block["note"])}</p>' if block["note"] else ""
            items = ("<ul>" + "".join(f"<li>{_inline(i)}</li>" for i in block["items"]) + "</ul>"
                     if block["items"] else "")
            html.append(f'<div class="job">{head}{note}{items}</div>')
        elif block["kind"] == "list":
            html.append("<ul>" + "".join(f"<li>{_inline(i)}</li>" for i in block["items"]) + "</ul>")
        else:
            html.append(f"<p>{_inline(block['text'])}</p>")
    return "".join(html)


def render_resume_html(md: str) -> str:
    """Build the two-column resume page."""
    doc = _parse_resume(md)

    side, main = [], []
    for section in doc["sections"]:
        target = side if section["heading"].strip().lower() in SIDEBAR_SECTIONS else main
        target.append(f'<section><h2>{_inline(section["heading"])}</h2>'
                      f'{_render_blocks(section["blocks"])}</section>')

    title = f'<p class="title">{_inline(doc["title"])}</p>' if doc["title"] else ""
    contact = f'<div class="contact">{_inline(doc["contact"])}</div>' if doc["contact"] else ""

    return (
        "<!doctype html><html><head><meta charset='utf-8'></head><body>"
        f'<div class="head"><div class="who"><h1 class="name">{_inline(doc["name"])}</h1>'
        f'{title}</div>{contact}</div><div class="rule"></div>'
        f'<div class="body"><div class="side">{"".join(side)}</div>'
        f'<div class="main">{"".join(main)}</div></div>'
        "</body></html>"
    )


# --------------------------------------------------------------------------
# Resume markdown -> single column, for upload
# --------------------------------------------------------------------------

# Sections whose list is a list of terms rather than of achievements. In one
# column a term per line wastes the page, and a comma-separated line is also
# the form skills parsers split most reliably.
INLINE_LIST_SECTIONS = {"skills", "technical skills", "core skills", "key skills"}


def _ats_blocks(blocks: list, inline_lists: bool) -> str:
    html = []
    for block in blocks:
        if block["kind"] == "job":
            dates = (f'<span class="jobdate"> | {_ats_inline(block["dates"])}</span>'
                     if block["dates"] else "")
            head = f'<p class="jobhead"><strong>{_ats_inline(block["title"])}</strong>{dates}</p>'
            note = f'<p class="jobnote">{_ats_inline(block["note"])}</p>' if block["note"] else ""
            items = "".join(f"<li>\u2022 {_ats_inline(i)}</li>" for i in block["items"])
            html.append(f'<div class="job">{head}{note}'
                        + (f"<ul>{items}</ul>" if items else "") + "</div>")
        elif block["kind"] == "list" and inline_lists:
            html.append(f"<p>{', '.join(_ats_inline(i) for i in block['items'])}</p>")
        elif block["kind"] == "list":
            html.append("<ul>" + "".join(f"<li>\u2022 {_ats_inline(i)}</li>"
                                         for i in block["items"]) + "</ul>")
        else:
            html.append(f"<p>{_ats_inline(block['text'])}</p>")
    return "".join(html)


def _ats_inline(text: str) -> str:
    return _inline(text, show_urls=True)


def render_resume_ats_html(md: str) -> str:
    """Build the single-column page, sections in the order they were written."""
    doc = _parse_resume(md)
    sections = "".join(
        f'<h2>{_ats_inline(s["heading"])}</h2>'
        + _ats_blocks(s["blocks"], s["heading"].strip().lower() in INLINE_LIST_SECTIONS)
        for s in doc["sections"]
    )
    title = f'<p class="title">{_ats_inline(doc["title"])}</p>' if doc["title"] else ""
    contact = f'<p class="contact">{_ats_inline(doc["contact"])}</p>' if doc["contact"] else ""
    return (
        "<!doctype html><html><head><meta charset='utf-8'></head><body>"
        f'<h1 class="name">{_ats_inline(doc["name"])}</h1>{title}{contact}{sections}'
        "</body></html>"
    )


def _plain_runs(paragraph, text: str, bold: bool = False, italic: bool = False) -> None:
    """Add markdown text to a Word paragraph as runs: **bold** kept, a link
    shown as its address (see _inline), everything else as written."""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", lambda m: _visible_url(m.group(2)), text)
    for i, piece in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not piece:
            continue
        run = paragraph.add_run(piece)
        run.bold = bold or (i % 2 == 1)
        run.italic = italic


def write_resume_docx(md: str, path: str, pt: float = 10.0) -> None:
    """The single-column resume as a Word document.

    Built from the same parse as the PDF, in the same order. Everything is body
    text: the section header and footer stay empty, because a parser that
    skips them would lose whatever was put there. No tables and no text boxes.
    Role lines put the dates after a right-aligned tab, which reads as one line
    of text to a parser and as a right-aligned date to a person.

    `pt` is the size the PDF fitted at, so the two copies match. Word lays the
    page out itself when it opens the file, so page count is not measured here.
    """
    from docx import Document
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Inches, Pt, RGBColor

    doc = _parse_resume(md)
    word = Document()
    section = word.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(0.5))
    usable = section.page_width - section.left_margin - section.right_margin

    normal = word.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(pt)
    normal.paragraph_format.space_after = Pt(pt * 0.3)
    normal.paragraph_format.space_before = Pt(0)

    heading = word.styles["Heading 1"]
    heading.font.name = "Arial"
    heading.font.size = Pt(pt * 1.1)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    heading.paragraph_format.space_before = Pt(pt * 0.9)
    heading.paragraph_format.space_after = Pt(pt * 0.3)

    name = word.add_paragraph()
    run = name.add_run(doc["name"])
    run.bold = True
    run.font.size = Pt(pt * 2.0)
    if doc["title"]:
        _plain_runs(word.add_paragraph(), doc["title"].strip("*"), bold=True)
    if doc["contact"]:
        _plain_runs(word.add_paragraph(), doc["contact"])

    for part in doc["sections"]:
        word.add_heading(part["heading"], level=1)
        inline_lists = part["heading"].strip().lower() in INLINE_LIST_SECTIONS
        for block in part["blocks"]:
            if block["kind"] == "job":
                head = word.add_paragraph()
                head.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
                head.paragraph_format.keep_with_next = True
                _plain_runs(head, block["title"], bold=True)
                if block["dates"]:
                    head.add_run("\t" + block["dates"])
                if block["note"]:
                    _plain_runs(word.add_paragraph(), block["note"], italic=True)
                for item in block["items"]:
                    _plain_runs(word.add_paragraph(style="List Bullet"), item)
            elif block["kind"] == "list" and inline_lists:
                _plain_runs(word.add_paragraph(), ", ".join(block["items"]))
            elif block["kind"] == "list":
                for item in block["items"]:
                    _plain_runs(word.add_paragraph(style="List Bullet"), item)
            else:
                _plain_runs(word.add_paragraph(), block["text"])

    word.save(path)


def _html(markdown_text: str) -> str:
    body = markdown.markdown(
        markdown_text,
        # nl2br keeps a role's location/date line and its title-history line on
        # separate lines instead of collapsing them into one paragraph.
        extensions=["tables", "sane_lists", "nl2br"],
    )
    return f"<!doctype html><html><head><meta charset='utf-8'></head><body>{body}</body></html>"


def write_pdf(markdown_text, path, style="report"):
    """Render an internal report. No length constraint."""
    HTML(string=_html(markdown_text)).write_pdf(path, stylesheets=[CSS(string=REPORT_CSS)])


def fit_pdf(markdown_text, path, style, max_pages=1):
    """Render a document, tightening the typography until it fits.

    Returns (pages, body_pt). Content is never altered — if even the tightest
    setting overruns, the document is left there and the caller is told.
    """
    builders = {"resume": render_resume_ats_html, "resume_designed": render_resume_html}
    html = builders.get(style, _html)(markdown_text)
    ladder, build = LADDERS[style], BUILDERS[style]
    for pt, margin in ladder:
        HTML(string=html).write_pdf(path, stylesheets=[CSS(string=build(pt, margin))])
        if page_count(path) <= max_pages:
            return page_count(path), pt
    return page_count(path), ladder[-1][0]

def page_count(path: str) -> int:
    """How many pages a rendered PDF actually came to."""
    return len(PdfReader(path).pages)
